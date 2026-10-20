"""Render reports/report.md into a polished reports/MLOps_Assignment1_Report.docx.

A small, purpose-built Markdown -> docx renderer (headings, paragraphs, bullet lists, pipe
tables, fenced code blocks, images, bold/inline-code spans, links and blockquotes) - good
enough for the structured Markdown this project's report.md actually contains, without pulling
in a heavyweight pandoc dependency.

Usage (from project root, with requirements.txt installed):
    python scripts/generate_report.py
"""
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from config import PROJECT_ROOT, REPORTS_DIR

REPORT_MD = REPORTS_DIR / "report.md"
REPORT_DOCX = REPORTS_DIR / "MLOps_Assignment1_Report.docx"

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))")
IMAGE_LINE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
TABLE_ROW = re.compile(r"^\|(.+)\|\s*$")
TABLE_SEPARATOR = re.compile(r"^\|[\s:|-]+\|\s*$")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_runs(paragraph, text: str) -> None:
    for chunk in INLINE_PATTERN.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif chunk.startswith("[") and "](" in chunk:
            label, href = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", chunk).groups()
            if href.startswith("#"):
                paragraph.add_run(label).italic = True  # internal anchor: keep the label only
            else:
                add_hyperlink(paragraph, label, href)
        else:
            paragraph.add_run(chunk)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.paragraphs[0].text = ""
            add_inline_runs(cell.paragraphs[0], cell_text.strip())
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def add_image(doc: Document, alt: str, relative_path: str) -> None:
    image_path = (REPORTS_DIR / relative_path).resolve()
    if not image_path.exists():
        p = doc.add_paragraph()
        run = p.add_run(f"[missing figure: {relative_path}]")
        run.italic = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        return
    doc.add_picture(str(image_path), width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(alt)
    run.italic = True
    run.font.size = Pt(9)


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    p._p.get_or_add_pPr().append(shading)
    for i, line in enumerate(lines):
        run = p.add_run(line if i == 0 else "\n" + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def render(markdown_text: str, doc: Document) -> None:
    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue

        image_match = IMAGE_LINE.match(line)
        if image_match:
            add_image(doc, image_match.group(1), image_match.group(2))
            i += 1
            continue

        if TABLE_ROW.match(line) and i + 1 < len(lines) and TABLE_SEPARATOR.match(lines[i + 1]):
            table_lines = [line]
            i += 2  # skip the header separator row
            while i < len(lines) and TABLE_ROW.match(lines[i]):
                table_lines.append(lines[i])
                i += 1
            rows = [[cell for cell in row.strip("|").split("|")] for row in table_lines]
            add_table(doc, rows)
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=min(level, 4))
            i += 1
            continue

        if line.strip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line.strip()[2:])
            i += 1
            continue

        if re.match(r"^\d+\.\s", line.strip()):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s", "", line.strip()))
            i += 1
            continue

        if line.strip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(line.strip().lstrip(">").strip())
            run.italic = True
            i += 1
            continue

        if line.strip() == "---":
            doc.add_paragraph("_" * 60)
            i += 1
            continue

        if line.strip():
            p = doc.add_paragraph()
            add_inline_runs(p, line.strip())

        i += 1


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    render(REPORT_MD.read_text(), doc)

    doc.save(REPORT_DOCX)
    print(f"Wrote {REPORT_DOCX.relative_to(PROJECT_ROOT)}")


if __name__ == "__main__":
    main()
